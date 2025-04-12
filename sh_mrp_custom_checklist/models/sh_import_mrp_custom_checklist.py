#-*- coding: utf-8 -*-
# Copyright (C) Softhealer Technologies.
 
from odoo import models, fields , api
import base64
import io
import csv
import xlrd
import openpyxl
 
class Sh_import_mrp_custom_checklist(models.TransientModel):
    _name = 'sh.import.mrp.custom.checklist'
    _description = 'Import MRP Custom Checklist'
    
    import_type = fields.Selection([('csv','CSV File') , ('excel','Excel File')] ,
                                    string="Import File Type", required=True, default='csv')
    
    company = fields.Many2one('res.company',string='Company' , default = lambda self:self.env.company)
    
    file = fields.Binary(string='Upload Your File', required=True)
    file_name = fields.Char(string='Filename')
    
    
    def action_import_file(self):
        self.ensure_one()
        data = base64.b64decode(self.file)
 
        if self.import_type == 'csv':
            self._import_csv(data)
        elif self.import_type == 'excel':
            self._import_excel(data)
 
        return {'type': 'ir.actions.act_window_close'}
 
    def _import_csv(self, data):
        file_stream = io.StringIO(data.decode("utf-8"))
        reader = csv.DictReader(file_stream)
        for row in reader:
            self.env['sh.mrp.checklist'].create({
                'name': row.get('name'),
                'description': row.get('description'),
            })
        return{
            self.env["bus.bus"]._sendone(
                self.env.user.partner_id,
                "simple_notification",
                {
                    "type": "success",
                    "title": "Imported",
                    "message": "Data importated successfully",
                },
                )
        }
    
     
    def _import_excel(self, data):
        file_stream = io.BytesIO(data)
        wb = openpyxl.load_workbook(file_stream)
        sheet = wb.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = dict(zip(headers, row))
            self.env['sh.mrp.checklist'].create({
                'name': row_data.get('name'),
                'description': row_data.get('description'),
            })
        return{
            self.env["bus.bus"]._sendone(
                self.env.user.partner_id,
                "simple_notification",
                {
                    "type": "success",
                    "title": "Imported",
                    "message": "Data importated successfully",
                },
                )
        }      